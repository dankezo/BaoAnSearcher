# -*- coding: utf-8 -*-
"""Rebuild Technical Survey Excel + clean architecture PNG for Genextech."""
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Side, PatternFill, Font, NamedStyle
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont

OUT_XLSX = Path(r"D:\BaoAnSearcher\outputs\BaoAn_Pharma_Technical_Survey_FILLED.xlsx")
OUT_PNG = Path(r"D:\BaoAnSearcher\outputs\BaoAn_Current_Architecture.png")
DESK = Path.home() / "Desktop" / "BaoAn_Pharma_Technical_Survey_FILLED.xlsx"


def load_fonts():
    try:
        return {
            "title": ImageFont.truetype("segoeuib.ttf", 26),
            "h": ImageFont.truetype("segoeuib.ttf", 15),
            "b": ImageFont.truetype("segoeui.ttf", 13),
            "s": ImageFont.truetype("segoeui.ttf", 11),
        }
    except OSError:
        f = ImageFont.load_default()
        return {"title": f, "h": f, "b": f, "s": f}


def draw_box(d, xy, fill, outline, title, lines, fonts, title_fill="#111827"):
    x1, y1, x2, y2 = xy
    d.rounded_rectangle(xy, radius=12, fill=fill, outline=outline, width=2)
    d.text((x1 + 14, y1 + 10), title, fill=title_fill, font=fonts["h"])
    yy = y1 + 36
    for line in lines:
        d.text((x1 + 14, yy), line, fill="#374151", font=fonts["s"])
        yy += 16


def h_arrow(d, x1, y, x2, label="", fonts=None, color="#4B5563"):
    """Horizontal arrow left→right or right→left, clean head, optional label above."""
    d.line([(x1, y), (x2, y)], fill=color, width=2)
    if x2 >= x1:
        d.polygon([(x2, y), (x2 - 10, y - 5), (x2 - 10, y + 5)], fill=color)
        mid = (x1 + x2) // 2
    else:
        d.polygon([(x2, y), (x2 + 10, y - 5), (x2 + 10, y + 5)], fill=color)
        mid = (x2 + x1) // 2
    if label and fonts:
        d.text((mid - len(label) * 3, y - 18), label, fill=color, font=fonts["s"])


def v_arrow(d, x, y1, y2, label="", fonts=None, color="#4B5563"):
    """Vertical arrow top→bottom or bottom→top."""
    d.line([(x, y1), (x, y2)], fill=color, width=2)
    if y2 >= y1:
        d.polygon([(x, y2), (x - 5, y2 - 10), (x + 5, y2 - 10)], fill=color)
        mid = (y1 + y2) // 2
    else:
        d.polygon([(x, y2), (x - 5, y2 + 10), (x + 5, y2 + 10)], fill=color)
        mid = (y2 + y1) // 2
    if label and fonts:
        d.text((x + 8, mid - 6), label, fill=color, font=fonts["s"])


def elbow_arrow(d, points, label="", fonts=None, color="#4B5563"):
    """Polyline with arrow head at end. points = [(x,y), ...]"""
    for i in range(len(points) - 1):
        d.line([points[i], points[i + 1]], fill=color, width=2)
    x0, y0 = points[-2]
    x1, y1 = points[-1]
    # head based on last segment direction
    if abs(x1 - x0) >= abs(y1 - y0):
        if x1 >= x0:
            d.polygon([(x1, y1), (x1 - 10, y1 - 5), (x1 - 10, y1 + 5)], fill=color)
        else:
            d.polygon([(x1, y1), (x1 + 10, y1 - 5), (x1 + 10, y1 + 5)], fill=color)
    else:
        if y1 >= y0:
            d.polygon([(x1, y1), (x1 - 5, y1 - 10), (x1 + 5, y1 - 10)], fill=color)
        else:
            d.polygon([(x1, y1), (x1 - 5, y1 + 10), (x1 + 5, y1 + 10)], fill=color)
    if label and fonts:
        # place near midpoint of first segment
        mx = (points[0][0] + points[1][0]) // 2
        my = (points[0][1] + points[1][1]) // 2
        d.text((mx + 6, my - 14), label, fill=color, font=fonts["s"])


def make_diagram():
    fonts = load_fonts()
    w, h = 1480, 980
    img = Image.new("RGB", (w, h), "#F8FAFC")
    d = ImageDraw.Draw(img)

    d.text((40, 20), "Bao An Pharma — Current Architecture (2026-09-26)", fill="#0F172A", font=fonts["title"])
    d.text(
        (40, 54),
        "Corporate site (public) + BaoAn Searcher (internal) · Cloudflare · Vercel · Supabase · Turso · Local crawl",
        fill="#64748B",
        font=fonts["b"],
    )

    # Layout boxes (x1,y1,x2,y2)
    users = (40, 100, 260, 210)
    cf = (320, 100, 620, 210)
    web_corp = (680, 90, 1080, 200)
    web_app = (1120, 90, 1440, 200)
    api = (900, 260, 1280, 400)
    supabase = (40, 280, 400, 450)
    turso = (40, 500, 420, 690)
    local = (500, 500, 920, 720)
    notes = (980, 500, 1440, 720)

    draw_box(
        d,
        users,
        "#EEF2FF",
        "#6366F1",
        "Users",
        ["Public: partners / visitors", "Internal: @baoanpharma.com", "Login: Outlook / Email"],
        fonts,
    )
    draw_box(
        d,
        cf,
        "#ECFDF5",
        "#10B981",
        "1. Cloudflare CDN / DNS / WAF",
        ["Zone: baoanpharma.com (proxied)", "TLS + WAF + security headers", "Worker: security-txt (minor)"],
        fonts,
    )
    draw_box(
        d,
        web_corp,
        "#FFF7ED",
        "#F97316",
        "2. Vercel — Corporate Web",
        ["baoanpharma.com  ·  project baoan-b2b", "Marketing / product catalog", "Region edge: HKG"],
        fonts,
    )
    draw_box(
        d,
        web_app,
        "#FFF7ED",
        "#EA580C",
        "3. Vercel — Searcher SPA",
        ["app.baoanpharma.com", "project baoan-searcher", "Vite + React (auth-gated)"],
        fonts,
    )
    draw_box(
        d,
        api,
        "#FEF3C7",
        "#D97706",
        "4. Vercel Serverless API",
        [
            "Node.js  /api/tender/*",
            "Auth: Supabase JWT Bearer",
            "Query: Turso libSQL (DAV/MSC/VSS)",
            "Managed Linux (serverless)",
        ],
        fonts,
    )
    draw_box(
        d,
        supabase,
        "#F0F9FF",
        "#0EA5E9",
        "5. Supabase Auth + Postgres",
        [
            "Project: baoan-auth",
            "Region: ap-northeast-1 (Tokyo)",
            "PostgreSQL 17 · Auth / RLS",
            "Azure AD + email allowlist",
            "DNS: *.supabase.co",
        ],
        fonts,
    )
    draw_box(
        d,
        turso,
        "#FDF4FF",
        "#A855F7",
        "6. Turso libSQL (business data)",
        [
            "DB: baoan-searcher",
            "Region: aws-ap-northeast-1",
            "~1.2 GB  ·  DAV 55k · VSS 747k",
            "MSC prices 10k · tenders 10k",
            "Private (server token only)",
        ],
        fonts,
    )
    draw_box(
        d,
        local,
        "#F1F5F9",
        "#64748B",
        "7. Local Crawl / Sync Workstation",
        [
            "Host: DANNY-PHAN · Windows 11 Home",
            "CPU: Ryzen 7 H 255 (8c/16t) · RAM 32 GB",
            "Disk: C ~500 GB · D ~450 GB",
            "Stack: FastAPI + SQLite + browser crawl",
            "Sources: DAV · MSC · VSS  →  sync cloud",
            "Hours: 8/5 + scheduled jobs · Private",
        ],
        fonts,
    )
    draw_box(
        d,
        notes,
        "#FFFFFF",
        "#94A3B8",
        "Survey notes",
        [
            "No dedicated production VM for public web/API",
            "Web/API = Vercel serverless (shared capacity)",
            "DB = managed SaaS (Supabase + Turso)",
            "vCPU/RAM in sheet = plan / equivalent estimate",
            "GPU: not used",
            "Company: CP Duoc My Pham Bao An",
            "MST 0103984595 · Ha Dong, Ha Noi",
        ],
        fonts,
    )

    # Clean single arrows — never cross through other boxes
    # Users → Cloudflare (corporate public path)
    h_arrow(d, users[2], 145, cf[0], "HTTPS", fonts)
    # Cloudflare → Corporate Vercel
    h_arrow(d, cf[2], 145, web_corp[0], "proxy", fonts)
    # Users → Searcher SPA (direct DNS→Vercel, not CF-proxied)
    # route under CF/corporate row, into left side of Searcher SPA
    elbow_arrow(
        d,
        [(150, 210), (150, 235), (1140, 235), (1140, 200)],
        "HTTPS direct → app (no CF proxy)",
        fonts,
        color="#6366F1",
    )
    # Searcher SPA → API (down into top of API)
    v_arrow(d, 1200, web_app[3], api[1], "UI→API", fonts)
    # API → Supabase Auth (left, clear of other boxes)
    elbow_arrow(d, [(900, 310), (420, 310)], "Auth JWT", fonts, color="#0EA5E9")
    # API → Turso: go LEFT then DOWN in the GAP between Turso and Local (x≈460)
    # so the line never crosses Local crawl box
    elbow_arrow(
        d,
        [(900, 380), (460, 380), (460, 560), (420, 560)],
        "Query data",
        fonts,
        color="#A855F7",
    )
    # Local → Turso Sync (ONE horizontal arrow only, mid-height)
    h_arrow(d, local[0], 620, turso[2], "sync", fonts, color="#64748B")

    # Footer legend
    d.text(
        (40, 760),
        "Traffic notes: baoanpharma.com = Cloudflare-proxied → Vercel.  app.baoanpharma.com = DNS CNAME → Vercel (not CF proxy).",
        fill="#475569",
        font=fonts["s"],
    )
    d.text(
        (40, 786),
        "Data path: Local crawl workstation syncs gov sources → Turso. Searcher SPA calls Serverless API → Turso (auth via Supabase).",
        fill="#475569",
        font=fonts["s"],
    )

    OUT_PNG.parent.mkdir(parents=True, exist_ok=True)
    img.save(OUT_PNG, "PNG")
    print("png", OUT_PNG)


# Columns match Genextech sheet exactly (B..Q)
# No | Server Name/Function | Qty | Application/Server Group | Component | OS | vCPU | RAM | CPU% | RAM% | GPU | Storage | Env | Hours | Access | IP/DNS
ROWS = [
    {
        "no": 1,
        "name": "CDN / DNS / WAF",
        "qty": 1,
        "group": "Corporate Web",
        "component": "Cloudflare CDN + DNS + WAF (+ security-txt Worker)",
        "os": "SaaS (Managed)",
        "vcpu": "Shared",
        "ram": "Shared",
        "cpu_pct": 10,
        "ram_pct": 15,
        "gpu": "N/A",
        "storage": 5,
        "env": "Production",
        "hours": "24/7",
        "access": "Public (via Internet)",
        "dns": "DNS: baoanpharma.com (Cloudflare proxied)",
    },
    {
        "no": 2,
        "name": "Web / App server — Corporate site",
        "qty": 1,
        "group": "Corporate Web / B2B",
        "component": "Vercel Static Hosting (project baoan-b2b) — marketing & product catalog",
        "os": "Vercel Managed Linux",
        "vcpu": 2,
        "ram": 2,
        "cpu_pct": 15,
        "ram_pct": 25,
        "gpu": "N/A",
        "storage": 5,
        "env": "Production",
        "hours": "24/7",
        "access": "Public (via Internet)",
        "dns": "DNS: baoanpharma.com",
    },
    {
        "no": 3,
        "name": "Web / App server — BaoAn Searcher SPA",
        "qty": 1,
        "group": "BaoAn Searcher",
        "component": "Vercel Static Hosting (project baoan-searcher) — Vite + React SPA",
        "os": "Vercel Managed Linux",
        "vcpu": 2,
        "ram": 2,
        "cpu_pct": 20,
        "ram_pct": 30,
        "gpu": "N/A",
        "storage": 5,
        "env": "Production",
        "hours": "24/7",
        "access": "Private (auth @baoanpharma.com)",
        "dns": "DNS: app.baoanpharma.com",
    },
    {
        "no": 4,
        "name": "API server — Searcher Serverless",
        "qty": 1,
        "group": "BaoAn Searcher",
        "component": "Vercel Serverless Node.js (/api/tender/search|metrics|meta) → Turso",
        "os": "Vercel Managed Linux",
        "vcpu": 2,
        "ram": 4,
        "cpu_pct": 25,
        "ram_pct": 40,
        "gpu": "N/A",
        "storage": 2,
        "env": "Production",
        "hours": "24/7",
        "access": "Private (JWT via Supabase Auth)",
        "dns": "DNS: app.baoanpharma.com/api/*",
    },
    {
        "no": 5,
        "name": "Database server — Auth / Metadata",
        "qty": 1,
        "group": "Identity & Auth",
        "component": "Supabase PostgreSQL 17 (project baoan-auth) + Auth/RLS + Azure AD",
        "os": "Managed Linux (Supabase)",
        "vcpu": 2,
        "ram": 1,
        "cpu_pct": 20,
        "ram_pct": 45,
        "gpu": "N/A",
        "storage": 8,
        "env": "Production",
        "hours": "24/7",
        "access": "Private (service + authenticated clients)",
        "dns": "DNS: *.supabase.co (ap-northeast-1)",
    },
    {
        "no": 6,
        "name": "Database server — Business Data",
        "qty": 1,
        "group": "BaoAn Searcher Data",
        "component": "Turso libSQL baoan-searcher — DAV/MSC/VSS (~1.2 GB live)",
        "os": "Managed Linux (Turso/AWS)",
        "vcpu": 2,
        "ram": 4,
        "cpu_pct": 35,
        "ram_pct": 55,
        "gpu": "N/A",
        "storage": 10,
        "env": "Production",
        "hours": "24/7",
        "access": "Private (server-side token only)",
        "dns": "DNS: baoan-searcher-*.turso.io (aws-ap-northeast-1)",
    },
    {
        "no": 7,
        "name": "Crawl / Sync Workstation",
        "qty": 1,
        "group": "Data Ops / Admin",
        "component": "Python FastAPI (uvicorn) + SQLite + browser crawl (DAV/MSC/VSS) → sync cloud",
        "os": "Windows 11 Home",
        "vcpu": 16,
        "ram": 32,
        "cpu_pct": 30,
        "ram_pct": 50,
        "gpu": "N/A",
        "storage": 200,
        "env": "Ops / Development",
        "hours": "8/5 (+ scheduled jobs)",
        "access": "Private (operator LAN/VPN)",
        "dns": "Host: DANNY-PHAN (local IP)",
    },
]


def build_workbook():
    wb = openpyxl.Workbook()

    # --- Sheet 1: paste-ready data (exact Genextech columns) ---
    ws = wb.active
    ws.title = "Technical Survey FILL"
    headers = [
        "No",
        "Server Name/Function",
        "Quantity",
        "Application/ Server Group",
        "Component",
        "OS",
        "vCPU",
        "RAM (GB)",
        "CPU Utilization (%)",
        "RAM Utilization (%)",
        "GPU VRAM (if needed)",
        "Total Storage Size (GB)",
        "Environment/ Purpose",
        "Operating Hours",
        "Public/Private Access (user)",
        "IP/DNS access (user)",
    ]
    header_fill = PatternFill("solid", fgColor="F97316")
    header_font = Font(bold=True, color="FFFFFF")
    wrap = Alignment(wrap_text=True, vertical="center")
    center = Alignment(wrap_text=True, vertical="center", horizontal="center")
    thin = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    for i, h in enumerate(headers, 1):
        cell = ws.cell(1, i, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = thin

    keys = [
        "no",
        "name",
        "qty",
        "group",
        "component",
        "os",
        "vcpu",
        "ram",
        "cpu_pct",
        "ram_pct",
        "gpu",
        "storage",
        "env",
        "hours",
        "access",
        "dns",
    ]
    for r_i, row in enumerate(ROWS, 2):
        for c_i, k in enumerate(keys, 1):
            cell = ws.cell(r_i, c_i, row[k])
            cell.border = thin
            cell.alignment = center if c_i in (1, 3, 7, 8, 9, 10, 12) else wrap
        ws.row_dimensions[r_i].height = 42

    widths = [5, 38, 10, 22, 55, 24, 10, 10, 12, 12, 14, 14, 16, 18, 34, 42]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:P{1 + len(ROWS)}"

    # --- Sheet 2: How to paste ---
    guide = wb.create_sheet("HUONG DAN DIEN")
    guide["A1"] = "Cách điền vào Google Sheet Genextech (đúng cột)"
    guide["A1"].font = Font(bold=True, size=14)
    steps = [
        "",
        "1) Mở sheet 'Technical Survey FILL' trong file này.",
        "2) Bôi đen A2:P8 (7 dòng dữ liệu, KHÔNG copy header).",
        "3) Ctrl+C.",
        "4) Vào Google Sheet Genextech → click ô cột No dòng số 1 (ô B5 trong template gốc).",
        "5) Paste ĐẶC BIỆT: Edit → Paste special → Paste values only (hoặc Ctrl+Shift+V).",
        "   → Tránh lệch cột như lần trước (đừng paste vào giữa bảng).",
        "",
        "MAP CỘT (bắt buộc khớp trái → phải):",
        "  No | Server Name/Function | Quantity | Application/Server Group | Component | OS | vCPU | RAM | CPU% | RAM% | GPU | Storage | Env | Hours | Access | IP/DNS",
        "",
        "LỖI THƯỜNG GẶP (như screenshot của bạn):",
        "  - Cột 'Application/Server Group' bị TRỐNG",
        "  - Cột 'Server Name/Function' lại chứa text kiểu 'Corporate Web', 'Identity & Auth'",
        "  → Đó là dán lệch 1 cột: Group bị đẩy sang Name, Name bị mất.",
        "  → Xóa 7 dòng đã dán sai → dán lại theo bước trên.",
        "",
        "Giữ nguyên 2 dòng Eg. (ví dụ xám) — KHÔNG ghi đè.",
        "",
        "Ảnh architecture: chèn file BaoAn_Current_Architecture.png vào vùng Picture (C17).",
        "  Desktop / outputs đều có bản PNG đã sửa mũi tên.",
        "",
        "Ghi chú PaaS: Cloudflare/Vercel/Supabase/Turso không phải VM riêng — vCPU/RAM là ước lượng tương đương plan.",
        "Máy #7 (DANNY-PHAN) là số thật: 16 vCPU / 32 GB RAM / Windows 11.",
    ]
    for i, line in enumerate(steps, 2):
        guide[f"A{i}"] = line
        guide[f"A{i}"].font = Font(size=11)
    guide.column_dimensions["A"].width = 120

    # --- Sheet 3: with diagram for reference ---
    arch = wb.create_sheet("Architecture PNG")
    arch["A1"] = "Current Architect Diagram (embed PNG below)"
    arch["A1"].font = Font(bold=True, size=12)
    arch["A2"] = (
        "Internet → Cloudflare (baoanpharma.com only) → Vercel corporate. "
        "app.baoanpharma.com → Vercel SPA/API trực tiếp. "
        "API → Supabase Auth + Turso data. Local workstation sync → Turso."
    )
    arch["A2"].alignment = wrap
    arch.merge_cells("A2:G2")
    arch.row_dimensions[2].height = 40
    img = XLImage(str(OUT_PNG))
    img.width = 960
    img.height = 640
    arch.add_image(img, "A4")

    wb.save(OUT_XLSX)
    desk_xlsx = Path.home() / "Desktop" / "BaoAn_Technical_Survey_FILLED_v2.xlsx"
    desk_png = Path.home() / "Desktop" / "BaoAn_Current_Architecture_v2.png"
    for dest, src in ((desk_xlsx, OUT_XLSX), (desk_png, OUT_PNG)):
        try:
            dest.write_bytes(src.read_bytes())
            print("desktop", dest)
        except PermissionError:
            alt = dest.with_name(dest.stem + "_new" + dest.suffix)
            alt.write_bytes(src.read_bytes())
            print("desktop locked, wrote", alt)
    try:
        DESK.write_bytes(OUT_XLSX.read_bytes())
        print("desktop", DESK)
    except PermissionError:
        print("skip locked", DESK)
    print("xlsx", OUT_XLSX)


if __name__ == "__main__":
    make_diagram()
    build_workbook()
