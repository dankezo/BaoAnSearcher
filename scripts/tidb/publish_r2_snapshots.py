# -*- coding: utf-8 -*-
"""Stream TiDB datasets into CSV snapshots and publish them to Cloudflare R2.

Full exports use TiFlash, so they do not compete with TiKV index lookups. The
CSV is streamed to a temporary file, then uploaded multipart: no million-row
list is kept in memory.

  python scripts/tidb/publish_r2_snapshots.py --yes-remote
  python scripts/tidb/publish_r2_snapshots.py --yes-remote --only MSC_PRICE
"""
from __future__ import annotations

import argparse
import csv
import os
import sys
import tempfile
import json
import subprocess
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from connect import connect, load_env, require_config

DATASETS = {
    "DAV": ("dav_drugs", "dav.csv"),
    "MSC_PRICE": ("msc_prices", "msc-price.csv"),
    "MSC_BID": ("msc_tenders", "msc-bid.csv"),
    "VSS": ("vss_bids", "vss.csv"),
}
CHUNK = 2_000


def required(name: str) -> str:
    value = (os.environ.get(name) or "").strip()
    if not value:
        raise SystemExit(f"Thiếu {name}. Xem .env.example.")
    return value


def r2_client():
    try:
        import boto3
        from botocore.config import Config
    except ImportError as exc:
        raise SystemExit("Thiếu boto3. Cài server requirements trước khi xuất R2.") from exc
    return boto3.client(
        "s3", endpoint_url=required("R2_ENDPOINT"), region_name="auto",
        aws_access_key_id=required("R2_ACCESS_KEY_ID"),
        aws_secret_access_key=required("R2_SECRET_ACCESS_KEY"),
        config=Config(signature_version="s3v4", retries={"max_attempts": 5, "mode": "standard"}),
    )


def public_url(key: str) -> str:
    return f"{required('R2_PUBLIC_BASE_URL').rstrip('/')}/{key}"


def columns_of(conn, table: str) -> list[str]:
    with conn.cursor() as cur:
        cur.execute(
            "SELECT COLUMN_NAME FROM information_schema.columns "
            "WHERE table_schema=DATABASE() AND table_name=%s ORDER BY ORDINAL_POSITION", (table,),
        )
        return [str(row[0]) for row in cur.fetchall()]


def write_csv(conn, code: str, table: str) -> tuple[Path, int]:
    import pymysql

    columns = columns_of(conn, table)
    if not columns:
        raise RuntimeError(f"Không thấy bảng {table}.")
    fd, raw_path = tempfile.mkstemp(prefix=f"baoan-{code.lower()}-", suffix=".csv")
    os.close(fd)
    path = Path(raw_path)
    count = 0
    quoted = ", ".join(f"`{column}`" for column in columns)
    sql = f"SELECT /*+ READ_FROM_STORAGE(TIFLASH[{table}]) */ {quoted} FROM `{table}`"
    try:
        with conn.cursor(pymysql.cursors.SSCursor) as cur, path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle, lineterminator="\n")
            writer.writerow(columns)
            cur.execute(sql)
            while rows := cur.fetchmany(CHUNK):
                writer.writerows(rows)
                count += len(rows)
    except Exception:
        path.unlink(missing_ok=True)
        raise
    return path, count


def update_registry(conn, code: str, count: int, url: str, size_mb: float) -> None:
    # Building a full workbook can outlast TiDB's idle connection timeout.
    conn.ping(reconnect=True)
    with conn.cursor() as cur:
        cur.execute(
            "UPDATE data_registry_meta SET total_records=%s, status='healthy', "
            "last_synced_at=CURRENT_TIMESTAMP, r2_download_url=%s, file_size_mb=%s "
            "WHERE dataset_code=%s", (count, url, size_mb, code),
        )
    conn.commit()


def write_excel(csv_path: Path, code: str) -> Path:
    """Build once beside CSV; users then download directly from R2."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font, PatternFill
    columns = json.loads((ROOT/'web/src/dataExportColumns.json').read_text(encoding='utf-8'))[code]
    path = csv_path.with_suffix('.xlsx')
    workbook = Workbook(write_only=True)
    has_summary = any(column.get('sum') for column in columns)
    data_rows_per_sheet = 1048574
    count = 0
    sheet = None
    from openpyxl.utils import get_column_letter
    def append_summary(target):
        if not has_summary:
            return
        summary = []
        for index, column in enumerate(columns,1):
            value = f'=SUM({get_column_letter(index)}3:{get_column_letter(index)}1048576)' if column.get('sum') else ('Tổng' if index == 1 else None)
            cell = WriteOnlyCell(target,value=value)
            cell.font = Font(bold=True,color='234E3D')
            cell.fill = PatternFill('solid',fgColor='E8F0EB')
            summary.append(cell)
        target.append(summary)
    with csv_path.open(encoding='utf-8-sig',newline='') as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            if count % data_rows_per_sheet == 0:
                sheet = workbook.create_sheet(f'{code}_{count//data_rows_per_sheet+1}')
                sheet.freeze_panes = 'A3' if has_summary else 'A2'
                headers = []
                for column in columns:
                    cell = WriteOnlyCell(sheet,value=column['label'])
                    cell.font = Font(bold=True,color='FFFFFF')
                    cell.fill = PatternFill('solid',fgColor='155E75')
                    headers.append(cell)
                sheet.append(headers)
                append_summary(sheet)
            numeric_values = {}
            from decimal import Decimal
            for column in columns:
                if column.get('numeric'):
                    value = row.get(column.get('cloudKey',column['key']), '')
                    if value != '':
                        value = Decimal(value)
                        if not value.is_finite():
                            raise ValueError('Snapshot có số không hợp lệ.')
                    else:
                        value = None
                    numeric_values[column['key']] = value
            if code == 'VSS' and 'thanhtien' in numeric_values:
                price, quantity = numeric_values.get('gia'), numeric_values.get('soluong')
                numeric_values['thanhtien'] = price * quantity if price is not None and quantity is not None else None
            elif code == 'MSC_PRICE' and 'total_value' in numeric_values:
                price, quantity = numeric_values.get('unit_price'), numeric_values.get('quantity')
                numeric_values['total_value'] = price * quantity if price is not None and quantity is not None else None
            cells = []
            for column in columns:
                value = numeric_values[column['key']] if column.get('numeric') else row.get(column.get('cloudKey',column['key']), '')
                if isinstance(value,str):
                    value = ILLEGAL_CHARACTERS_RE.sub('',value)
                cell = WriteOnlyCell(sheet,value=value)
                if isinstance(value,str):
                    cell.data_type = 's'
                    cell.number_format = '@'
                elif column.get('numeric') and value is not None:
                    cell.number_format = '#,##0' if value == value.to_integral_value() else '#,##0.############'
                cells.append(cell)
            sheet.append(cells)
            count += 1
    if sheet is None:
        sheet=workbook.create_sheet(code)
        sheet.freeze_panes='A3' if has_summary else 'A2'
        sheet.append([column['label'] for column in columns])
        append_summary(sheet)
    for index, worksheet in enumerate(workbook.worksheets):
        worksheet.auto_filter.ref=f'A1:{get_column_letter(len(columns))}{min(count-index*data_rows_per_sheet,data_rows_per_sheet)+(2 if has_summary else 1)}'
    workbook.save(path)
    return path


def publish_one(conn, client, bucket: str, code: str, excel: bool = False) -> None:
    table, filename = DATASETS[code]
    print(f"export {code} from TiFlash…", flush=True)
    path, count = write_csv(conn, code, table)
    print(f'exported {code}: {count:,} rows; preparing downloads',flush=True)
    key = f"snapshots/{filename}"
    size_mb = round(path.stat().st_size / 1024 / 1024, 2)
    try:
        if excel:
            excel_path = write_excel(path,code)
            try:
                excel_name = Path(filename).with_suffix('.xlsx').name
                excel_key = f'snapshots/{excel_name}'
                client.upload_file(str(excel_path),bucket,excel_key,ExtraArgs={
                    'ContentType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    'ContentDisposition':f'attachment; filename="{excel_name}"',
                    'CacheControl':'public, max-age=300',
                })
                client.head_object(Bucket=bucket,Key=excel_key)
                print(f'published {code} Excel',flush=True)
            finally:
                excel_path.unlink(missing_ok=True)
        client.upload_file(
            str(path), bucket, key,
            ExtraArgs={
                "ContentType": "text/csv; charset=utf-8",
                "ContentDisposition": f'attachment; filename="{filename}"',
                "CacheControl": "public, max-age=300",
            },
        )
        client.head_object(Bucket=bucket, Key=key)
        update_registry(conn, code, count, public_url(key), size_mb)
    finally:
        path.unlink(missing_ok=True)
    print(f"published {code}: {count:,} rows, {size_mb:,.2f} MB", flush=True)


def publish_analytics() -> None:
    bucket = required('R2_BUCKET')
    client = r2_client()
    with tempfile.TemporaryDirectory(prefix='baoan-analytics-') as directory:
        result = subprocess.run(['node', str(ROOT / 'scripts' / 'build_analytics_snapshots.mjs'), '--build', '--out', directory], cwd=ROOT, capture_output=True, text=True, encoding='utf-8', timeout=7200, creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
        if result.returncode:
            raise RuntimeError('Analytics snapshot build failed; previous manifest retained.')
        folder=Path(directory)
        manifest=json.loads((folder/'manifest.json').read_text(encoding='utf-8'))
        for item in manifest['entries']:
            path=folder/item['path']
            client.upload_file(str(path),bucket,'analytics/'+item['path'],ExtraArgs={'ContentType':'application/json; charset=utf-8','CacheControl':'public, max-age=31536000, immutable'})
            client.head_object(Bucket=bucket,Key='analytics/'+item['path'])
        # The manifest becomes visible only after every object has been verified.
        client.upload_file(str(folder/'manifest.json'),bucket,'analytics/manifest.json',ExtraArgs={'ContentType':'application/json; charset=utf-8','CacheControl':'public, max-age=60'})
        print(f"Published {len(manifest['entries'])} aggregate snapshots.",flush=True)


def main() -> None:
    parser = argparse.ArgumentParser(description="Publish TiDB CSV snapshots directly to Cloudflare R2.")
    parser.add_argument("--yes-remote", action="store_true", help="Export live TiDB data and upload to R2.")
    parser.add_argument("--only", choices=tuple(DATASETS), help="Publish one dataset.")
    parser.add_argument("--analytics", action="store_true", help="Publish versioned aggregate analytics instead of raw CSV.")
    parser.add_argument('--excel', action='store_true', help='Publish a compressed Excel snapshot alongside each CSV.')
    args = parser.parse_args()
    if not args.yes_remote:
        print("No data exported. Pass --yes-remote to create and upload R2 snapshots.")
        return
    load_env()
    if args.analytics:
        publish_analytics()
        return
    bucket = required("R2_BUCKET")
    conn = connect(require_config())
    try:
        client = r2_client()
        for code in ([args.only] if args.only else DATASETS): publish_one(conn, client, bucket, code,excel=args.excel)
    finally:
        conn.close()


if __name__ == "__main__":
    main()
