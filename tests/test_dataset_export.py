from contextlib import closing
import csv
from datetime import datetime
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from server import dataset_export, vss


class DatasetExportTest(unittest.TestCase):
    def test_computed_line_values_and_non_sum_sheet_layout(self):
        with tempfile.TemporaryDirectory() as directory:
            folder=Path(directory)
            source=folder/'source.sqlite3'
            source.touch()
            cases=(
                ('VSS',vss,'VSS_DB',{'soluong':'0.5','gia':'4.25','thanhtien':'999','_analyticsNumbers':{'soluong':'0.5','gia':'4.25','thanhtien':'999'}}),
                ('MSC_PRICE',dataset_export.msc,'MSC_DB',{'unit_price':'1.25','quantity':'0.4','total_value':'999'}),
                ('DAV',dataset_export.dav,'DAV_DB',{'soDangKy':'001'}),
            )
            for code,module,source_key,row in cases:
                with self.subTest(code=code), patch.object(module,source_key,source), patch.object(dataset_export,'DATA_DIR',folder), patch.object(dataset_export,'source_rows',return_value=iter([row])):
                    path,_,_=dataset_export.create_export({'code':code,'format':'xlsx'})
                    with closing(load_workbook(path,data_only=False)) as workbook:
                        sheet=workbook[f'{code}_1']
                        self.assertEqual(sheet['A1'].value,dataset_export.COLUMNS[code][0]['label'])
                        if code=='VSS':
                            total_col=next(i for i,col in enumerate(dataset_export.COLUMNS[code],1) if col['key']=='thanhtien')
                            self.assertEqual(sheet.cell(3,total_col).value,2.125)
                        elif code=='MSC_PRICE':
                            total_col=next(i for i,col in enumerate(dataset_export.COLUMNS[code],1) if col['key']=='total_value')
                            self.assertEqual(sheet.cell(3,total_col).value,0.5)
                        else:
                            self.assertEqual(sheet.freeze_panes,'A2')
                            self.assertEqual(sheet['A2'].value,'001')

    def test_excel_removes_source_control_characters_and_preserves_literal_text(self):
        from scripts.tidb.publish_r2_snapshots import write_excel
        with tempfile.TemporaryDirectory() as directory:
            source=Path(directory)/'vss.csv'
            with source.open('w',encoding='utf-8-sig',newline='') as handle:
                writer=csv.DictWriter(handle,fieldnames=['sodk','ten','nhasx','soluong'])
                writer.writeheader()
                writer.writerow(dict(sodk='0590110425723',ten='=1+1',nhasx='\x08RV GROUP (S) PTE LTD',soluong='100.000'))
            with closing(load_workbook(write_excel(source,'VSS'),read_only=True)) as workbook:
                rows=list(workbook['VSS_1'].iter_rows())
                values=[cell.value for cell in rows[2]]
                self.assertIn('RV GROUP (S) PTE LTD',values)
                self.assertEqual(rows[1][7].value,'=SUM(H3:H1048576)')
                self.assertEqual(rows[2][1].value,'0590110425723')
                self.assertEqual(rows[2][1].data_type,'s')
                self.assertEqual(rows[2][2].value,'=1+1')
                self.assertEqual(rows[2][2].data_type,'s')
                self.assertEqual(rows[2][7].value,100)

    def test_snapshot_excel_computes_msc_total_without_changing_csv(self):
        from scripts.tidb.publish_r2_snapshots import write_excel
        columns=dataset_export.COLUMNS['MSC_PRICE']
        with tempfile.TemporaryDirectory() as directory:
            source=Path(directory)/'prices.csv'
            with source.open('w',encoding='utf-8-sig',newline='') as handle:
                writer=csv.DictWriter(handle,fieldnames=[column['key'] for column in columns])
                writer.writeheader()
                writer.writerow({'unit_price':'1.25','quantity':'0.4','total_value':'999'})
            excel=write_excel(source,'MSC_PRICE')
            with closing(load_workbook(excel,data_only=False)) as workbook:
                sheet=workbook['MSC_PRICE_1']
                total_col=next(i for i,column in enumerate(columns,1) if column['key']=='total_value')
                self.assertEqual(sheet.freeze_panes,'A3')
                self.assertTrue(sheet.cell(2,total_col).value.startswith('=SUM('))
                self.assertEqual(sheet.cell(3,total_col).value,0.5)
            with source.open(encoding='utf-8-sig',newline='') as handle:
                self.assertEqual(next(csv.DictReader(handle))['total_value'],'999')

    def test_full_export_preserves_numbers_identifiers_and_date_window(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            source = folder / 'vss.sqlite3'
            with patch.object(vss,'VSS_DB',source), patch.object(dataset_export,'DATA_DIR',folder):
                with closing(vss.connect()) as con, con:
                    for index in range(5104):
                        day = '2026-10-01 23:59:59.999' if index < 5102 else ('2026-10-02' if index == 5102 else '')
                        raw = dict(sodk='0590110425723',loai='Tân dược',tungay_hd=day,hoatchat='Lacidipin',ten='=1+1',nhasx='\x08RV GROUP (S) PTE LTD',dangbaoche='Viên nén',soluong='100.000',gia='4.200',thanhtien='420.000',_analyticsNumbers={'soluong':'100.000','gia':'4200.00','thanhtien':'420000.00'})
                        con.execute('INSERT INTO bids(fingerprint,raw,search,sodk,loai,tungay_hd,hoatchat) VALUES(?,?,?,?,?,?,?)',(str(index),json.dumps(raw),'lacidipin','0590110425723','Tân dược',day,'Lacidipin'))
                body=dict(code='VSS',filters={'sodk':'0590110425723'},start='2025-06-01',end='2026-10-01',format='xlsx')
                path,_,_ = dataset_export.create_export(body)
                with closing(load_workbook(path,data_only=False)) as workbook:
                    rows = list(workbook['VSS_1'].iter_rows())
                    self.assertEqual(len(rows),5104)
                    self.assertEqual(workbook['VSS_1'].freeze_panes,'A3')
                    self.assertEqual(rows[1][0].value,'Tổng')
                    self.assertEqual([rows[1][index].value for index in (7,8,9)],['=SUM(H3:H1048576)','=SUM(I3:I1048576)','=SUM(J3:J1048576)'])
                    self.assertTrue(rows[1][7].font.bold)
                    self.assertEqual(rows[1][7].fill.fgColor.rgb[-6:],'E8F0EB')
                    self.assertEqual(rows[2][1].value,'0590110425723')
                    self.assertEqual(rows[2][1].data_type,'s')
                    self.assertEqual(rows[2][2].value,'=1+1')
                    self.assertEqual(rows[2][2].data_type,'s')
                    self.assertIn('RV GROUP (S) PTE LTD',[cell.value for cell in rows[2]])
                    self.assertEqual(sum(row[7].value for row in rows[2:]),510200)
                    self.assertEqual(rows[2][7].number_format,'#,##0')
                    self.assertEqual(sum(row[9].value for row in rows[2:]),2142840000)
                    self.assertEqual(workbook['Bo_loc']['B4'].value,5102)
                self.assertIsInstance(rows[2][17].value,datetime)
                stamp=path.stat().st_mtime_ns
                self.assertEqual(dataset_export.create_export(body)[0].stat().st_mtime_ns,stamp)
                csv_path,_,_=dataset_export.create_export({**body,'format':'csv'})
                with csv_path.open(encoding='utf-8-sig',newline='') as handle:
                    rows=list(csv.reader(handle))
                    self.assertEqual(rows[1][2],"'=1+1")
                    self.assertEqual(len(rows),5103)
                with self.assertRaisesRegex(ValueError,'trước'):
                    dataset_export.create_export({**body,'start':'2027-01-01'})
                with self.assertRaises(ValueError):
                    dataset_export.create_export({**body,'end':'2026-02-30'})


if __name__ == '__main__':
    unittest.main()
