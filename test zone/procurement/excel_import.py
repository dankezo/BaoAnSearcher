"""Import DAV procurement exports without losing the original Excel observations.

Exact equal Excel rows are represented once, with every original row retained.
An Excel record is hidden in favour of API only for a unique, exact normalized
business-field match. Ambiguous matches stay visible. Future API saves reconcile
the same index automatically; no source records are deleted.
"""
import hashlib
import json
import re
import sys
import unicodedata
from collections import defaultdict
from datetime import datetime,date
from decimal import Decimal,InvalidOperation
from pathlib import Path
from core import ROOT,connect,save_records,normalize_price,now

HEADERS=['STT','Tên thuốc','Tên hoạt chất/ thành phần dược liệu','Nồng độ, hàm lượng','GĐKLH hoặc GPNK',
 'Đường dùng','Dạng bào chế','Hạn dùng (Tuổi thọ)','Tên cơ sở sản xuất','Nước sản xuất',
 'Quy cách đóng gói','Đơn vị tính','Số lượng','Đơn giá trúng thầu','Mã định danh NT trúng thầu',
 'Tên NT trúng thầu','Nhóm thuốc','Mã TBMT','Mã định danh CĐT','Tên CĐT','Hình thức LCNT',
 'Ngày đăng tải KQLCNT','Số quyết định','Ngày ban hành quyết định','Số nhà thầu tham dự','Địa điểm']
RAW_FIELDS=['tenThuoc','tenHoatChat','nongDo','gdklh_GPNK','duongDung','dangBaoChe','hanDung',
 'tenCoSoSanXuat','nuocSanXuat','quyCachDongGoi','donViTinh','soLuong','donGia','winningCode',
 'winningName','nhomThuoc','maTbmt','maCdt','tenCdtBmt','bidForm','ngayDangTaiKqlcnt',
 'soQuyetDinh','ngayBanHanhQuyetDinh','bidderCount','diaDiem']
MATCH_FIELDS=['tender_no','name','ingredient','strength','registration','route','dosage_form',
 'manufacturer','country','packaging','unit','quantity','unit_price','winner_code','winner',
 'buyer_code','buyer','group_name','published','decision','decision_date']

def norm(value):return ' '.join(unicodedata.normalize('NFC',str(value or '')).casefold().split())
def digest(value):return hashlib.sha256(json.dumps(value,ensure_ascii=False,sort_keys=True,default=str,separators=(',',':')).encode()).hexdigest()

def vn_number(value):
    if value in (None,''):return None
    if isinstance(value,(int,float)):return float(value)
    text=str(value).strip().replace('\u00a0','').replace(' ','')
    if not re.fullmatch(r'-?(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d+)?',text):
        raise ValueError('Số theo định dạng Việt Nam chưa hợp lệ: '+text)
    result=Decimal(text.replace('.','').replace(',','.'))
    return float(result)

def iso_date(value):
    if value in (None,''):return ''
    if isinstance(value,(datetime,date)):return value.strftime('%Y-%m-%dT00:00:00')
    return datetime.strptime(str(value).strip(),'%d/%m/%Y').strftime('%Y-%m-%dT00:00:00')

def fingerprint(obj):
    # Incomplete product identity must never cause automatic hiding.
    if any(obj.get(k) in (None,'') for k in ('tender_no','name','registration','unit','quantity','unit_price')):return None
    values=[]
    for field in MATCH_FIELDS:
        value=obj.get(field)
        if field in ('quantity','unit_price'):
            value=str(Decimal(str(value)).normalize()) if value is not None else ''
        elif field in ('published','decision_date'):value=str(value or '')[:10]
        elif field=='group_name':
            match=re.fullmatch(r'(?:nhóm\s*|n\s*)?([1-5])',norm(value))
            if match:value=match.group(1)
        values.append(norm(value))
    return digest(values)

def index_price(con,ident,obj,is_excel):
    old=con.execute('SELECT fingerprint FROM price_keys WHERE source_id=?',(ident,)).fetchone()
    fp=fingerprint(obj)
    con.execute('DELETE FROM price_keys WHERE source_id=?',(ident,))
    con.execute('DELETE FROM excel_matches WHERE excel_id=? OR api_id=?',(ident,ident))
    if fp:con.execute('INSERT INTO price_keys VALUES(?,?,?)',(ident,fp,int(is_excel)))
    for value in {x for x in (fp,old[0] if old else None) if x}:
        ids=con.execute('SELECT source_id FROM price_keys WHERE fingerprint=? AND is_excel=0 LIMIT 2',(value,)).fetchall()
        con.execute('DELETE FROM excel_matches WHERE excel_id IN (SELECT source_id FROM price_keys WHERE fingerprint=? AND is_excel=1)',(value,))
        if len(ids)==1:
            con.execute('INSERT OR REPLACE INTO excel_matches SELECT source_id,? FROM price_keys WHERE fingerprint=? AND is_excel=1',(ids[0][0],value))

def import_xlsx(path,db=None):
    sys.path.insert(0,str(ROOT/'vendor'))
    from openpyxl import load_workbook
    path=Path(path);file_hash=hashlib.sha256(path.read_bytes()).hexdigest()
    with connect(db) as con:
        previous=con.execute('SELECT row_count FROM excel_files WHERE file_hash=?',(file_hash,)).fetchone()
        if previous:return {'prices':0,'tenders':0,'skipped':previous[0],'already_imported':True}
    workbook=load_workbook(path,read_only=True,data_only=False)
    groups={};observations=[];exported='';sheets=0
    try:
        for sheet in workbook:
            rows=sheet.iter_rows();first=next(rows,None);header=next(rows,None)
            if not header or [norm(c.value) for c in header][:26]!=[norm(h) for h in HEADERS]:continue
            sheets+=1
            if first and first[0].value:exported=str(first[0].value)
            for row in rows:
                values=[cell.value for cell in row][:26]
                if not any(v not in (None,'') for v in values):continue
                if len(values)!=26:raise ValueError(f'{sheet.title}, dòng {row[0].row}: thiếu cột.')
                if any(cell.data_type in ('f','e') for cell in row):raise ValueError(f'{sheet.title}, dòng {row[0].row}: chứa công thức hoặc lỗi Excel, chưa nhập.')
                if not values[1] or not values[17]:raise ValueError(f'{sheet.title}, dòng {row[0].row}: thiếu tên thuốc hoặc mã TBMT.')
                ident='xlsx:'+digest(values[1:])
                if ident not in groups:
                    raw=dict(zip(RAW_FIELDS,values[1:]));raw.update(id=ident,tab='THUOC_TAN_DUOC')
                    try:
                        raw['donGia']=vn_number(values[13]);raw['soLuong']=vn_number(values[12])
                        raw['ngayDangTaiKqlcnt']=iso_date(values[21]);raw['ngayBanHanhQuyetDinh']=iso_date(values[23])
                    except (ValueError,InvalidOperation) as e:raise ValueError(f'{sheet.title}, dòng {row[0].row}: {e}') from e
                    # The export has no medicine category column. Only explicit labels identify it.
                    label=norm(values[16]);raw['medicines']='1' if label in ('bdg','bd','biệt dược','biệt dược gốc') else ''
                    raw['_excel']={'filename':path.name,'exported_at':exported,'note':'Không có ID dòng API; đối chiếu theo các trường dữ liệu.'}
                    groups[ident]=raw
                observations.append((file_hash,sheet.title,row[0].row,ident,json.dumps(values,ensure_ascii=False,default=str)))
    finally:workbook.close()
    if not sheets or not observations:raise ValueError('Không tìm thấy bảng Danh sách hàng hóa đúng 26 cột trong file.')
    counts=defaultdict(int)
    for obs in observations:counts[obs[3]]+=1
    for ident,count in counts.items():
        if count>1:groups[ident]['_excel']['note']=f'{count} dòng Excel giống nhau (trừ STT), gom hiển thị; chưa xác định là các phần thầu riêng hay dòng lặp. Bản gốc giữ đủ.'
    with connect(db) as con:
        # Populate the index for API records collected before this importer existed.
        for record in con.execute("SELECT source_id,normalized,raw FROM records WHERE kind='prices' AND source_id NOT IN (SELECT source_id FROM price_keys)").fetchall():
            index_price(con,record['source_id'],json.loads(record['normalized']),bool(json.loads(record['raw']).get('_excel')))
        existing={r[0] for r in con.execute("SELECT source_id FROM records WHERE kind='prices'")}
        fresh=[raw for ident,raw in groups.items() if ident not in existing]
        save_records(con,'prices',fresh)
        con.execute('INSERT INTO excel_files VALUES(?,?,?,?,?)',(file_hash,path.name,exported,now(),len(observations)))
        con.executemany('INSERT INTO excel_rows VALUES(?,?,?,?,?)',observations)
        matched=sum(con.execute('SELECT 1 FROM excel_matches WHERE excel_id=?',(ident,)).fetchone() is not None for ident in groups)
        visible_new=sum(raw['id'] not in {r[0] for r in con.execute('SELECT excel_id FROM excel_matches')} for raw in []) if False else 0
        linked={r[0] for r in con.execute('SELECT excel_id FROM excel_matches')}
        visible_new=sum(raw['id'] not in linked for raw in fresh)
        con.execute('INSERT INTO imports(imported_at,filename,prices,tenders,skipped) VALUES(?,?,?,?,?)',(now(),path.name,visible_new,0,0))
    return {'prices':visible_new,'tenders':0,'skipped':0,'source_rows':len(observations),
      'distinct_excel_rows':len(groups),'grouped_repeats':len(observations)-len(groups),'matched_api':matched}

def provenance(ident,db=None):
    with connect(db) as con:
        return [dict(r) for r in con.execute('''SELECT f.filename,e.sheet,count(*) row_count,
          min(e.row_no) first_row,max(e.row_no) last_row FROM excel_rows e
          JOIN excel_files f USING(file_hash) WHERE e.record_id=? OR e.record_id IN
          (SELECT excel_id FROM excel_matches WHERE api_id=?) GROUP BY f.file_hash,e.sheet''',(ident,ident))]
